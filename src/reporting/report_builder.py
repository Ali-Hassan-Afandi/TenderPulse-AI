from io import BytesIO
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER

def _s(x):return str(x if x is not None else "")
def build_docx(company,tender,result):
    d=Document();p=d.add_paragraph();r=p.add_run("TenderPulse AI — Bid Intelligence Report");r.bold=True;r.font.size=Pt(20)
    for k,v in [("Company",company.get("company_name")),("Tender",tender.get("title")),("Tender ID",tender.get("id")),("Category",tender.get("category")),("Procedure",tender.get("procedure")),("Deadline",tender.get("deadline")),("Readiness",result["eligibility"].get("readiness")),("Risk",result["risk"].get("risk_level"))]:d.add_paragraph(f"{k}: {_s(v)}")
    d.add_heading("Compliance Matrix",1)
    for x in result["compliance"]["matrix"]:d.add_paragraph(f"{x.get('status')} — {x.get('requirement')}: {x.get('evidence')}")
    d.add_heading("AI Analysis",1);d.add_paragraph(_s(result["document"]["ai"].get("summary")))
    d.add_paragraph("Decision support only. Verify all requirements against the official tender.")
    b=BytesIO();d.save(b);return b.getvalue()

def build_pdf(company,tender,result):
    b=BytesIO();styles=getSampleStyleSheet()
    title=ParagraphStyle("TPTitle",parent=styles["Title"],alignment=TA_CENTER,fontSize=18,leading=22)
    body=ParagraphStyle("TPBody",parent=styles["BodyText"],fontSize=9,leading=13,spaceAfter=6)
    story=[Paragraph("TenderPulse AI",title),Paragraph("Bid Intelligence Report",styles["Heading2"]),Spacer(1,8)]
    fields=[("Company",company.get("company_name")),("Tender",tender.get("title")),("Tender ID",tender.get("id")),("Category",tender.get("category")),("Procedure",tender.get("procedure")),("Deadline",tender.get("deadline")),("Readiness",f"{result['eligibility'].get('readiness',0)}%"),("Opportunity Fit",f"{result['eligibility'].get('fit',0)}%"),("Risk",result["risk"].get("risk_level"))]
    for k,v in fields:story.append(Paragraph(f"<b>{k}:</b> {_s(v)}",body))
    story+=[Spacer(1,8),Paragraph("Compliance Matrix",styles["Heading2"])]
    for x in result["compliance"]["matrix"]:
        story.append(Paragraph(f"<b>{_s(x.get('status'))}</b> — {_s(x.get('requirement'))}<br/>{_s(x.get('evidence'))}<br/><i>{_s(x.get('action'))}</i>",body))
    story+=[Spacer(1,8),Paragraph("AI Analysis",styles["Heading2"]),Paragraph(_s(result["document"]["ai"].get("summary")) or "No AI narrative available.",body)]
    for risk in result["risk"].get("risks",[]):story.append(Paragraph("• "+_s(risk),body))
    story.append(Spacer(1,8));story.append(Paragraph("<b>Important:</b> This report is decision support, not an award prediction. Verify the official procurement documents before submission.",body))
    SimpleDocTemplate(b,pagesize=A4,rightMargin=38,leftMargin=38,topMargin=38,bottomMargin=38,title="TenderPulse AI Report").build(story)
    return b.getvalue()
